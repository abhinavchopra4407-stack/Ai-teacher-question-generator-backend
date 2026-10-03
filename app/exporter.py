import io
from typing import List, Dict, Any, Optional
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def generate_docx_paper(
    title: str,
    subject: str,
    grade: str,
    board: str,
    school_name: str,
    teacher_name: str,
    total_marks: int,
    instructions: str,
    questions: List[Dict[str, Any]],
    answer_key: Optional[List[Dict[str, Any]]] = None
) -> bytes:
    """Generate professional DOCX question paper with optional Answer Key."""
    doc = docx.Document()
    
    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)
        
    # Header Title
    head_para = doc.add_paragraph()
    head_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_school = head_para.add_run(school_name.upper() + "\n")
    run_school.font.size = Pt(16)
    run_school.font.bold = True
    run_school.font.color.rgb = RGBColor(30, 58, 138) # Deep Blue
    
    run_exam = head_para.add_run(f"EXAMINATION QUESTION PAPER: {title.upper()}\n")
    run_exam.font.size = Pt(13)
    run_exam.font.bold = True
    
    # Metadata Table
    meta_table = doc.add_table(rows=2, cols=3)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False
    
    row1 = meta_table.rows[0].cells
    row1[0].text = f"Subject: {subject}"
    row1[1].text = f"Class/Grade: {grade}"
    row1[2].text = f"Total Marks: {total_marks}"
    
    row2 = meta_table.rows[1].cells
    row2[0].text = f"Board: {board}"
    row2[1].text = f"Teacher: {teacher_name or 'N/A'}"
    row2[2].text = "Time Allowed: 2 Hours"
    
    for row in meta_table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(10)
                    r.font.bold = True
                    
    doc.add_paragraph()
    
    # Student Info Box
    student_table = doc.add_table(rows=1, cols=2)
    st_cells = student_table.rows[0].cells
    st_cells[0].text = "Student Name: __________________________"
    st_cells[1].text = "Roll No / ID: ___________________"
    
    doc.add_paragraph()
    
    # Instructions
    if instructions:
        inst_p = doc.add_paragraph()
        r_inst_head = inst_p.add_run("General Instructions:\n")
        r_inst_head.font.bold = True
        r_inst_head.font.size = Pt(10)
        r_inst_body = inst_p.add_run(instructions)
        r_inst_body.font.italic = True
        r_inst_body.font.size = Pt(10)
        
    doc.add_paragraph("―" * 45).alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Group questions by section
    vs_q = [q for q in questions if "Very Short" in q.get("question_type", "")]
    s_q = [q for q in questions if "Short" in q.get("question_type", "") and "Very" not in q.get("question_type", "")]
    l_q = [q for q in questions if "Long" in q.get("question_type", "")]
    
    sections_list = [
        ("SECTION A: VERY SHORT ANSWER QUESTIONS", vs_q),
        ("SECTION B: SHORT ANSWER QUESTIONS", s_q),
        ("SECTION C: LONG ANSWER QUESTIONS", l_q)
    ]
    
    q_counter = 1
    for sec_title, q_group in sections_list:
        if not q_group:
            continue
        sec_p = doc.add_paragraph()
        sec_run = sec_p.add_run(f"\n{sec_title} ({len(q_group)} x {q_group[0].get('marks', 2)} = {len(q_group)*q_group[0].get('marks', 2)} Marks)")
        sec_run.font.bold = True
        sec_run.font.size = Pt(11)
        sec_run.font.color.rgb = RGBColor(30, 58, 138)
        
        for q in q_group:
            qp = doc.add_paragraph()
            r_num = qp.add_run(f"Q{q_counter}. ")
            r_num.font.bold = True
            
            qp.add_run(q.get("question_text", ""))
            
            r_marks = qp.add_run(f"  [{q.get('marks')} Marks]")
            r_marks.font.bold = True
            r_marks.font.color.rgb = RGBColor(100, 116, 139)
            
            q_counter += 1
            
    # Optional Answer Key Section
    if answer_key:
        doc.add_page_break()
        ak_head = doc.add_paragraph()
        ak_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r_ak = ak_head.add_run("OFFICIAL ANSWER KEY & MARKING SCHEME\n")
        r_ak.font.size = Pt(14)
        r_ak.font.bold = True
        r_ak.font.color.rgb = RGBColor(185, 28, 28) # Red Header
        
        ak_counter = 1
        for q in answer_key:
            ak_p = doc.add_paragraph()
            r_qnum = ak_p.add_run(f"Q{ak_counter} Answer ({q.get('question_type', '')} - {q.get('marks')} Marks):\n")
            r_qnum.font.bold = True
            
            if q.get("answer"):
                ak_p.add_run(f"Suggested Answer: {q.get('answer')}\n")
                
            if q.get("marking_points"):
                ak_p.add_run("Marking Scheme:\n")
                for pt in q.get("marking_points", []):
                    ak_p.add_run(f" • {pt}\n")
                    
            if q.get("expected_length"):
                r_len = ak_p.add_run(f"Expected Length: {q.get('expected_length')}\n")
                r_len.font.italic = True
                
            ak_counter += 1

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def generate_pdf_paper(
    title: str,
    subject: str,
    grade: str,
    board: str,
    school_name: str,
    teacher_name: str,
    total_marks: int,
    instructions: str,
    questions: List[Dict[str, Any]],
    answer_key: Optional[List[Dict[str, Any]]] = None
) -> bytes:
    """Generate printable PDF question paper with ReportLab."""
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )
    
    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'HeaderTitle',
        parent=styles['Normal'],
        fontSize=16,
        leading=20,
        alignment=1, # Center
        fontName='Helvetica-Bold',
        textColor=colors.HexColor('#1E3A8A')
    )
    
    sub_title_style = ParagraphStyle(
        'SubHeaderTitle',
        parent=styles['Normal'],
        fontSize=12,
        leading=16,
        alignment=1,
        fontName='Helvetica-Bold',
        textColor=colors.HexColor('#1F2937')
    )
    
    sec_style = ParagraphStyle(
        'SecTitle',
        parent=styles['Normal'],
        fontSize=11,
        leading=15,
        fontName='Helvetica-Bold',
        textColor=colors.HexColor('#1E3A8A'),
        spaceBefore=10,
        spaceAfter=6
    )
    
    q_style = ParagraphStyle(
        'QuestionText',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        fontName='Helvetica',
        spaceAfter=6
    )
    
    ans_style = ParagraphStyle(
        'AnswerText',
        parent=styles['Normal'],
        fontSize=9.5,
        leading=13,
        fontName='Helvetica-Oblique',
        textColor=colors.HexColor('#334155'),
        spaceAfter=4
    )

    story = []
    
    story.append(Paragraph(school_name.upper(), title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph(f"QUESTION PAPER: {title.upper()}", sub_title_style))
    story.append(Spacer(1, 8))
    
    # Metadata Table
    meta_data = [
        [f"Subject: {subject}", f"Class: {grade}", f"Total Marks: {total_marks}"],
        [f"Board: {board}", f"Teacher: {teacher_name or 'N/A'}", "Time: 2 Hours"]
    ]
    t = Table(meta_data, colWidths=[170, 170, 170])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#F8FAFC')),
        ('BOX', (0,0), (-1,-1), 1, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold'),
        ('FONTSIZE', (0,0), (-1,-1), 9),
        ('ALIGN', (0,0), (-1,-1), 'CENTER'),
    ]))
    story.append(t)
    story.append(Spacer(1, 10))
    
    # Student details line
    st_data = [["Student Name: _______________________", "Roll No: __________________"]]
    st_table = Table(st_data, colWidths=[260, 250])
    st_table.setStyle(TableStyle([('FONTSIZE', (0,0), (-1,-1), 9.5), ('FONTNAME', (0,0), (-1,-1), 'Helvetica-Bold')]))
    story.append(st_table)
    story.append(Spacer(1, 10))
    
    if instructions:
        story.append(Paragraph(f"<b>Instructions:</b> <i>{instructions}</i>", q_style))
        story.append(Spacer(1, 6))
        
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=10))
    
    vs_q = [q for q in questions if "Very Short" in q.get("question_type", "")]
    s_q = [q for q in questions if "Short" in q.get("question_type", "") and "Very" not in q.get("question_type", "")]
    l_q = [q for q in questions if "Long" in q.get("question_type", "")]
    
    sections_list = [
        ("SECTION A: VERY SHORT ANSWER QUESTIONS", vs_q),
        ("SECTION B: SHORT ANSWER QUESTIONS", s_q),
        ("SECTION C: LONG ANSWER QUESTIONS", l_q)
    ]
    
    q_counter = 1
    for sec_title, q_group in sections_list:
        if not q_group:
            continue
        story.append(Paragraph(f"{sec_title} ({len(q_group)} x {q_group[0].get('marks', 2)} = {len(q_group)*q_group[0].get('marks', 2)} Marks)", sec_style))
        for q in q_group:
            q_text = f"<b>Q{q_counter}.</b> {q.get('question_text')} &nbsp;&nbsp;<b>[{q.get('marks')} Marks]</b>"
            story.append(Paragraph(q_text, q_style))
            q_counter += 1
            
    # Optional Answer Key Page
    if answer_key:
        story.append(PageBreak())
        ak_title_style = ParagraphStyle(
            'AKTitle',
            parent=styles['Normal'],
            fontSize=15,
            alignment=1,
            fontName='Helvetica-Bold',
            textColor=colors.HexColor('#B91C1C')
        )
        story.append(Paragraph("OFFICIAL ANSWER KEY & MARKING SCHEME", ak_title_style))
        story.append(Spacer(1, 10))
        
        ak_cnt = 1
        for q in answer_key:
            head_txt = f"<b>Q{ak_cnt}. [{q.get('question_type')}] (Marks: {q.get('marks')})</b>"
            story.append(Paragraph(head_txt, q_style))
            if q.get("answer"):
                story.append(Paragraph(f"<b>Suggested Answer:</b> {q.get('answer')}", ans_style))
            if q.get("marking_points"):
                pts = "<br/>".join([f"• {pt}" for pt in q.get("marking_points", [])])
                story.append(Paragraph(f"<b>Marking Scheme:</b><br/>{pts}", ans_style))
            if q.get("expected_length"):
                story.append(Paragraph(f"<b>Expected Length:</b> {q.get('expected_length')}", ans_style))
            story.append(Spacer(1, 8))
            ak_cnt += 1

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()
